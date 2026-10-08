"""ATS Resume Checker - Streamlit + Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description,
and get an ATS score with concrete improvement suggestions.
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# Change this if Google retires/renames the model. You can also set it as
# a secret or environment variable called GEMINI_MODEL.
DEFAULT_MODEL = "gemini-2.5-flash"
MAX_RESUME_CHARS = 20000
MAX_JD_CHARS = 8000

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume honestly and strictly. \
Do not invent facts that are not in the resume.

Score the resume from 0 to 100 using these categories and maximum points:
- keywords: 25 (relevant skills/keywords; if a job description is given, match against it)
- formatting: 20 (ATS-friendly structure, standard headings, no tables/columns/graphics issues)
- experience: 20 (clear roles, dates, and impact)
- achievements: 15 (quantified, results-oriented bullet points)
- skills_education: 10 (clear skills and education sections)
- readability: 10 (concise, consistent, free of grammar/spelling errors)

The overall score MUST equal the sum of the category scores.

Respond ONLY with JSON in exactly this shape:
{
  "overall_score": <int 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "category_scores": {
    "keywords": <int 0-25>,
    "formatting": <int 0-20>,
    "experience": <int 0-20>,
    "achievements": <int 0-15>,
    "skills_education": <int 0-10>,
    "readability": <int 0-10>
  },
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {"priority": "High|Medium|Low", "section": "<resume section>", \
"issue": "<what is wrong>", "suggestion": "<specific fix>"}
  ],
  "rewrite_examples": [
    {"original": "<weak bullet from the resume>", "improved": "<stronger version>"}
  ]
}
Give 3-6 strengths, 3-6 weaknesses, up to 15 missing keywords (empty list if \
none), 5-10 improvements, and 2-4 rewrite examples."""

CATEGORY_MAX = {
    "keywords": 25,
    "formatting": 20,
    "experience": 20,
    "achievements": 15,
    "skills_education": 10,
    "readability": 10,
}
CATEGORY_LABELS = {
    "keywords": "Keywords",
    "formatting": "Formatting",
    "experience": "Experience",
    "achievements": "Achievements",
    "skills_education": "Skills & Education",
    "readability": "Readability",
}


# ----------------------------------------------------------------------------
# File reading
# ----------------------------------------------------------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Return plain text from a PDF, DOCX or TXT file's bytes."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(data))
        except Exception:
            raise ValueError("Could not open this PDF. The file may be corrupted.")
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------------------------------------------------------
# AI call + result handling
# ----------------------------------------------------------------------------
def _clamp(value, low, high):
    try:
        number = int(round(float(value)))
    except (TypeError, ValueError):
        number = 0
    return max(low, min(high, number))


def _str_list(value, limit=20):
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()][:limit]


def parse_response(raw: str) -> dict:
    """Parse and sanitize the model's JSON output."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start == -1 or end <= start:
            raise ValueError("The AI response was not valid JSON. Please try again.")
        data = json.loads(text[start : end + 1])
    if not isinstance(data, dict):
        raise ValueError("The AI response had an unexpected format. Please try again.")

    raw_cats = data.get("category_scores") or {}
    cats = {k: _clamp(raw_cats.get(k, 0), 0, mx) for k, mx in CATEGORY_MAX.items()}
    # Keep the overall score consistent with the categories when they exist.
    if any(cats.values()):
        overall = sum(cats.values())
    else:
        overall = _clamp(data.get("overall_score", 0), 0, 100)

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).title(),
                    "section": str(item.get("section", "General")),
                    "issue": str(item.get("issue", "")),
                    "suggestion": str(item.get("suggestion", "")),
                }
            )
    rewrites = []
    for item in data.get("rewrite_examples") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append(
                {"original": str(item["original"]), "improved": str(item["improved"])}
            )

    return {
        "overall_score": overall,
        "summary": str(data.get("summary", "")),
        "category_scores": cats,
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords"), 15),
        "improvements": improvements,
        "rewrite_examples": rewrites,
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    client = genai.Client(api_key=api_key)
    content = f"RESUME:\n{resume_text[:MAX_RESUME_CHARS]}\n\n"
    if job_description.strip():
        content += f"TARGET JOB DESCRIPTION:\n{job_description[:MAX_JD_CHARS]}\n"
    else:
        content += "No job description provided. Evaluate for general ATS-readiness.\n"

    response = client.models.generate_content(
        model=model,
        contents=content,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            response_mime_type="application/json",
            temperature=0.2,
        ),
    )
    return parse_response(response.text)


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def get_secret(name: str, default: str = "") -> str:
    """Read from Streamlit secrets first, then environment variables."""
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass
    return os.environ.get(name, default)


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 60:
        return "Good - needs some polish"
    if score >= 40:
        return "Fair - needs work"
    return "Poor - major improvements needed"


def show_results(result: dict) -> None:
    score = result["overall_score"]
    left, right = st.columns([1, 2])
    with left:
        st.metric("ATS Score", f"{score} / 100")
        st.caption(score_label(score))
        st.progress(score / 100)
    with right:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.subheader("Score breakdown")
    cols = st.columns(3)
    for i, (key, maximum) in enumerate(CATEGORY_MAX.items()):
        value = result["category_scores"][key]
        with cols[i % 3]:
            st.metric(CATEGORY_LABELS[key], f"{value} / {maximum}")
            st.progress(value / maximum)

    col_a, col_b = st.columns(2)
    with col_a:
        st.subheader("Strengths")
        for s in result["strengths"] or ["None listed."]:
            st.markdown(f"- {s}")
    with col_b:
        st.subheader("Weaknesses")
        for w in result["weaknesses"] or ["None listed."]:
            st.markdown(f"- {w}")

    if result["missing_keywords"]:
        st.subheader("Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.subheader("Suggested improvements")
    order = {"High": 0, "Medium": 1, "Low": 2}
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    for item in sorted(result["improvements"], key=lambda x: order.get(x["priority"], 1)):
        icon = icons.get(item["priority"], "🟠")
        with st.expander(f"{icon} {item['priority']} - {item['section']}"):
            st.markdown(f"**Issue:** {item['issue']}")
            st.markdown(f"**Fix:** {item['suggestion']}")

    if result["rewrite_examples"]:
        st.subheader("Rewrite examples")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex['original']}")
            st.markdown(f"**After:** {ex['improved']}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.write("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_secret("GEMINI_API_KEY")
    model = get_secret("GEMINI_MODEL", DEFAULT_MODEL)
    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password")
            st.caption("Get a free key at https://aistudio.google.com/apikey")
        else:
            st.success("API key loaded.")
        model = st.text_input("Model", value=model)

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)", height=150
    )

    if st.button("Analyze resume", type="primary"):
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return
        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.name, uploaded.getvalue())
            if len(text) < 100:
                st.error(
                    "Could not read enough text. If your PDF is a scanned image, "
                    "export a text-based PDF or upload a DOCX instead."
                )
                return
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, job_description, api_key, model)
            st.session_state["result"] = result
        except ValueError as err:
            st.error(str(err))
            return
        except Exception as err:  # network, quota, bad key, etc.
            st.error(f"Something went wrong: {err}")
            return

    if "result" in st.session_state:
        show_results(st.session_state["result"])

    st.caption(
        "Scores are AI estimates, not the output of a real ATS. Use them as guidance."
    )


if __name__ == "__main__":
    main()
